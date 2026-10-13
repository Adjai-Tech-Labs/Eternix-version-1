import os
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from eternix.core.models import Task, ToolDefinition, ToolResult, ToolSideEffect
from eternix.core.task_store import TaskStore
from eternix.tools.registry import tool_registry


async def doc_read_pdf_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    if not path_str:
        return ToolResult(
            tool_id="doc.read_pdf",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        import fitz
        doc = fitz.open(path_str)
        pages = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            pages.append({"page": page_num + 1, "text": text, "chars": len(text)})
        doc.close()

        return ToolResult(
            tool_id="doc.read_pdf",
            task_id=task.id,
            success=True,
            output={"path": path_str, "pages": pages, "total_pages": len(pages)},
        )
    except Exception as e:
        return ToolResult(
            tool_id="doc.read_pdf",
            task_id=task.id,
            success=False,
            error=f"Failed to read PDF: {e}",
        )


async def doc_read_docx_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    if not path_str:
        return ToolResult(
            tool_id="doc.read_docx",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        doc = Document(path_str)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = []
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                table_data.append([cell.text for cell in row.cells])
            tables.append(table_data)

        return ToolResult(
            tool_id="doc.read_docx",
            task_id=task.id,
            success=True,
            output={"path": path_str, "paragraphs": paragraphs, "tables": tables},
        )
    except Exception as e:
        return ToolResult(
            tool_id="doc.read_docx",
            task_id=task.id,
            success=False,
            error=f"Failed to read DOCX: {e}",
        )


async def doc_read_xlsx_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    if not path_str:
        return ToolResult(
            tool_id="doc.read_xlsx",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        import openpyxl
        wb = openpyxl.load_workbook(path_str, read_only=True, data_only=True)
        sheets = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append(list(row))
            sheets[sheet_name] = {"rows": rows, "dimensions": ws.dimensions}
        wb.close()

        return ToolResult(
            tool_id="doc.read_xlsx",
            task_id=task.id,
            success=True,
            output={"path": path_str, "sheets": sheets},
        )
    except Exception as e:
        return ToolResult(
            tool_id="doc.read_xlsx",
            task_id=task.id,
            success=False,
            error=f"Failed to read XLSX: {e}",
        )


async def doc_read_pptx_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    if not path_str:
        return ToolResult(
            tool_id="doc.read_pptx",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        prs = Presentation(path_str)
        slides = []
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        texts.append(paragraph.text)
            slides.append({"slide": i + 1, "texts": texts})

        return ToolResult(
            tool_id="doc.read_pptx",
            task_id=task.id,
            success=True,
            output={"path": path_str, "slides": slides, "total_slides": len(slides)},
        )
    except Exception as e:
        return ToolResult(
            tool_id="doc.read_pptx",
            task_id=task.id,
            success=False,
            error=f"Failed to read PPTX: {e}",
        )


async def doc_ocr_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    lang = input_data.get("lang", "eng")

    if not path_str:
        return ToolResult(
            tool_id="doc.ocr",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        import pytesseract
        from pdf2image import convert_from_path

        if path_str.lower().endswith(".pdf"):
            images = convert_from_path(path_str)
            results = []
            for i, img in enumerate(images):
                text = pytesseract.image_to_string(img, lang=lang)
                results.append({"page": i + 1, "text": text})
            return ToolResult(
                tool_id="doc.ocr",
                task_id=task.id,
                success=True,
                output={"path": path_str, "pages": results},
            )
        else:
            from PIL import Image
            img = Image.open(path_str)
            text = pytesseract.image_to_string(img, lang=lang)
            return ToolResult(
                tool_id="doc.ocr",
                task_id=task.id,
                success=True,
                output={"path": path_str, "text": text},
            )
    except Exception as e:
        return ToolResult(
            tool_id="doc.ocr",
            task_id=task.id,
            success=False,
            error=f"OCR failed: {e}",
        )


async def doc_create_docx_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    title = input_data.get("title", "Document")
    content = input_data.get("content", "")
    paragraphs = input_data.get("paragraphs", [])
    tables = input_data.get("tables", [])

    if not path_str:
        return ToolResult(
            tool_id="doc.create_docx",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        doc = Document()
        doc.add_heading(title, 0)

        if content:
            doc.add_paragraph(content)

        for p in paragraphs:
            doc.add_paragraph(p)

        for table_data in tables:
            if table_data:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                for i, row in enumerate(table_data):
                    for j, cell in enumerate(row):
                        table.cell(i, j).text = str(cell)

        doc.save(path_str)
        return ToolResult(
            tool_id="doc.create_docx",
            task_id=task.id,
            success=True,
            output={"path": path_str, "size": os.path.getsize(path_str)},
        )
    except Exception as e:
        return ToolResult(
            tool_id="doc.create_docx",
            task_id=task.id,
            success=False,
            error=f"Failed to create DOCX: {e}",
        )


async def doc_create_xlsx_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    sheets = input_data.get("sheets", {})

    if not path_str:
        return ToolResult(
            tool_id="doc.create_xlsx",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        wb = Workbook()
        default_sheet = wb.active
        default_sheet.title = "Sheet1"

        for sheet_name, sheet_data in sheets.items():
            if sheet_name == "Sheet1":
                ws = default_sheet
            else:
                ws = wb.create_sheet(sheet_name)

            rows = sheet_data.get("rows", [])
            for row_data in rows:
                ws.append(row_data)

        wb.save(path_str)
        return ToolResult(
            tool_id="doc.create_xlsx",
            task_id=task.id,
            success=True,
            output={"path": path_str, "size": os.path.getsize(path_str)},
        )
    except Exception as e:
        return ToolResult(
            tool_id="doc.create_xlsx",
            task_id=task.id,
            success=False,
            error=f"Failed to create XLSX: {e}",
        )


async def doc_create_pptx_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    title = input_data.get("title", "Presentation")
    slides = input_data.get("slides", [])

    if not path_str:
        return ToolResult(
            tool_id="doc.create_pptx",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        prs = Presentation()
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = title

        for slide_data in slides:
            layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
            slide = prs.slides.add_slide(layout)
            if slide.shapes.title:
                slide.shapes.title.text = slide_data.get("title", "")
            if len(slide.placeholders) > 1:
                body = slide.placeholders[1].text_frame
                body.text = slide_data.get("content", "")

        prs.save(path_str)
        return ToolResult(
            tool_id="doc.create_pptx",
            task_id=task.id,
            success=True,
            output={"path": path_str, "size": os.path.getsize(path_str)},
        )
    except Exception as e:
        return ToolResult(
            tool_id="doc.create_pptx",
            task_id=task.id,
            success=False,
            error=f"Failed to create PPTX: {e}",
        )


async def doc_create_pdf_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    path_str = input_data.get("path", "")
    title = input_data.get("title", "Report")
    content = input_data.get("content", "")
    sections = input_data.get("sections", [])

    if not path_str:
        return ToolResult(
            tool_id="doc.create_pdf",
            task_id=task.id,
            success=False,
            error="No path specified",
        )

    try:
        c = canvas.Canvas(path_str, pagesize=letter)
        _width, height = letter

        c.setFont("Helvetica-Bold", 24)
        c.drawString(72, height - 72, title)

        y = height - 120
        c.setFont("Helvetica", 12)

        if content:
            for line in content.split("\n"):
                c.drawString(72, y, line[:100])
                y -= 18
                if y < 72:
                    c.showPage()
                    y = height - 72

        for section in sections:
            if y < 120:
                c.showPage()
                y = height - 72
            c.setFont("Helvetica-Bold", 14)
            c.drawString(72, y, section.get("heading", ""))
            y -= 24
            c.setFont("Helvetica", 12)
            for line in section.get("content", "").split("\n"):
                c.drawString(72, y, line[:100])
                y -= 18
                if y < 72:
                    c.showPage()
                    y = height - 72

        c.save()
        return ToolResult(
            tool_id="doc.create_pdf",
            task_id=task.id,
            success=True,
            output={"path": path_str, "size": os.path.getsize(path_str)},
        )
    except Exception as e:
        return ToolResult(
            tool_id="doc.create_pdf",
            task_id=task.id,
            success=False,
            error=f"Failed to create PDF: {e}",
        )


async def doc_convert_handler(input_data: dict[str, Any], task: Task, task_store: TaskStore) -> ToolResult:
    src = input_data.get("src", "")
    dst = input_data.get("dst", "")

    if not src or not dst:
        return ToolResult(
            tool_id="doc.convert",
            task_id=task.id,
            success=False,
            error="Source and destination paths required",
        )

    try:
        import asyncio
        proc = await asyncio.create_subprocess_exec(
            "libreoffice",
            "--headless",
            "--convert-to",
            Path(dst).suffix[1:],
            "--outdir",
            str(Path(dst).parent),
            src,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        _stdout, stderr = await proc.communicate()
        if proc.returncode == 0 and os.path.exists(dst):
            return ToolResult(
                tool_id="doc.convert",
                task_id=task.id,
                success=True,
                output={"src": src, "dst": dst, "size": os.path.getsize(dst)},
            )
        else:
            stderr_text = stderr.decode("utf-8", errors="replace") if stderr else ""
            return ToolResult(
                tool_id="doc.convert",
                task_id=task.id,
                success=False,
                error=f"Conversion failed: {stderr_text}",
            )
    except Exception as e:
        return ToolResult(
            tool_id="doc.convert",
            task_id=task.id,
            success=False,
            error=f"Conversion failed: {e}",
        )


def register_document_tools() -> None:
    tool_registry.register(
        ToolDefinition(
            id="doc.read_pdf",
            name="Read PDF",
            description="Extract text and metadata from PDF files",
            input_schema={"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=60,
            side_effect=ToolSideEffect.READ_ONLY,
            tags=["document", "pdf", "read"],
        ),
        doc_read_pdf_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.read_docx",
            name="Read DOCX",
            description="Extract text and tables from DOCX files",
            input_schema={"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=30,
            side_effect=ToolSideEffect.READ_ONLY,
            tags=["document", "docx", "read"],
        ),
        doc_read_docx_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.read_xlsx",
            name="Read XLSX",
            description="Extract data from Excel spreadsheets",
            input_schema={"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=30,
            side_effect=ToolSideEffect.READ_ONLY,
            tags=["document", "xlsx", "read"],
        ),
        doc_read_xlsx_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.read_pptx",
            name="Read PPTX",
            description="Extract text from PowerPoint presentations",
            input_schema={"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=30,
            side_effect=ToolSideEffect.READ_ONLY,
            tags=["document", "pptx", "read"],
        ),
        doc_read_pptx_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.ocr",
            name="OCR Document",
            description="Extract text from images or scanned PDFs using OCR",
            input_schema={"type": "object", "properties": {"path": {"type": "string"}, "lang": {"type": "string", "default": "eng"}}, "required": ["path"]},
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=120,
            side_effect=ToolSideEffect.READ_ONLY,
            tags=["document", "ocr", "read"],
        ),
        doc_ocr_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.create_docx",
            name="Create DOCX",
            description="Create a Word document",
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string"},
                    "content": {"type": "string"},
                    "paragraphs": {"type": "array", "items": {"type": "string"}},
                    "tables": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}},
                },
                "required": ["path"],
            },
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=30,
            side_effect=ToolSideEffect.MODIFIES_FILES,
            tags=["document", "docx", "write"],
        ),
        doc_create_docx_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.create_xlsx",
            name="Create XLSX",
            description="Create an Excel spreadsheet",
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "sheets": {
                        "type": "object",
                        "additionalProperties": {
                            "type": "object",
                            "properties": {"rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}}},
                        },
                    },
                },
                "required": ["path"],
            },
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=30,
            side_effect=ToolSideEffect.MODIFIES_FILES,
            tags=["document", "xlsx", "write"],
        ),
        doc_create_xlsx_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.create_pptx",
            name="Create PPTX",
            description="Create a PowerPoint presentation",
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string"},
                    "slides": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {"title": {"type": "string"}, "content": {"type": "string"}},
                        },
                    },
                },
                "required": ["path"],
            },
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=30,
            side_effect=ToolSideEffect.MODIFIES_FILES,
            tags=["document", "pptx", "write"],
        ),
        doc_create_pptx_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.create_pdf",
            name="Create PDF",
            description="Create a PDF report using ReportLab",
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string"},
                    "content": {"type": "string"},
                    "sections": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {"heading": {"type": "string"}, "content": {"type": "string"}},
                        },
                    },
                },
                "required": ["path"],
            },
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=30,
            side_effect=ToolSideEffect.MODIFIES_FILES,
            tags=["document", "pdf", "write"],
        ),
        doc_create_pdf_handler,
    )

    tool_registry.register(
        ToolDefinition(
            id="doc.convert",
            name="Convert Document",
            description="Convert documents using LibreOffice",
            input_schema={"type": "object", "properties": {"src": {"type": "string"}, "dst": {"type": "string"}}, "required": ["src", "dst"]},
            output_schema={"type": "object"},
            execution_backend="python",
            timeout=120,
            side_effect=ToolSideEffect.MODIFIES_FILES,
            tags=["document", "convert"],
        ),
        doc_convert_handler,
    )